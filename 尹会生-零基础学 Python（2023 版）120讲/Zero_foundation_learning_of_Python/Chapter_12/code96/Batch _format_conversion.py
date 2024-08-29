# NOTE: 将word中的table提取，并存在一个新的excel文件中
import openpyxl
from docx import Document

doc = Document('./simple.docx')
table = doc.tables

wb = openpyxl.Workbook()
ws = wb.active

data_row = []
data_col = []
for t in table:
    for row in t.rows:
        for col in row.cells:
            data_col.append(col.text)
        data_row.append(data_col)
        data_col = []

for row in data_row:
    ws.append(row)

wb.save('simple.xlsx')

# NOTE: 总结
#  不同格式转换需要使用不同的库
#  格式转换一般需要打开多个文件，操作完成后需要将文件关闭，否则会导致内存占用过高
